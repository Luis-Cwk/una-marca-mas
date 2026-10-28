import sys

path = r"C:\Users\petra\agenteX\hf26-ofrenda\DesaparecidosMXFichas.xlsx"

try:
    import openpyxl
except ImportError:
    print("FALTA openpyxl")
    sys.exit(1)

wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("hojas:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("---- hoja:", name, "filas:", ws.max_row, "columnas:", ws.max_column)
    rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    header = next(rows)
    print("encabezados:", [str(h)[:40] if h is not None else None for h in header])

    # llenado por columna, sin imprimir valores
    filled = [0] * len(header)
    total = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        total += 1
        for i, v in enumerate(row[: len(filled)]):
            if v not in (None, ""):
                filled[i] += 1
    for h, n in zip(header, filled):
        pct = (100.0 * n / total) if total else 0
        print(f"  {str(h)[:40]!r}: {n}/{total} ({pct:.0f}%)")
